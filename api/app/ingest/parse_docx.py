"""DOCX extraction.

The cleanest path in the whole pipeline: Word carries real heading styles, so
the section hierarchy is read rather than guessed. Tables are flattened with
their header row repeated per data row, so a retrieved row still says what its
columns mean.

DOCX has no page concept until it is rendered, so `page` stays None and
citations fall back to the heading path.
"""

from __future__ import annotations

import io

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingest import chunk as chunker
from app.ingest.parse_pdf import ParseError
from app.schemas import ParsedDocument


def _heading_level(paragraph: Paragraph) -> int | None:
    name = (paragraph.style.name or "").lower() if paragraph.style else ""
    if name.startswith("heading"):
        tail = name.replace("heading", "").strip()
        if tail.isdigit():
            return min(int(tail), 6)
        return 1
    if name in ("title", "subtitle"):
        return 1
    return None


def _table_to_lines(table: Table) -> list[str]:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    header, *body = rows
    if not body:
        return [" | ".join(header)]
    lines = []
    for row in body:
        pairs = [
            f"{header[i] or f'kolom {i + 1}'}: {value}"
            for i, value in enumerate(row)
            if value
        ]
        if pairs:
            lines.append("; ".join(pairs))
    return lines


def _iter_body(document: Document):
    """Yield paragraphs and tables in document order."""
    body = document.element.body
    para_map = {p._element: p for p in document.paragraphs}
    table_map = {t._element: t for t in document.tables}
    for child in body.iterchildren():
        if child in para_map:
            yield para_map[child]
        elif child in table_map:
            yield table_map[child]


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        document = Document(io.BytesIO(data))
    except Exception:
        raise ParseError("DOCX tidak bisa dibuka atau rusak.")

    lines: list[str] = []
    for item in _iter_body(document):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                lines.append("")
                continue
            level = _heading_level(item)
            lines.append(("#" * level + " " + text) if level else text)
        else:
            lines.extend(_table_to_lines(item))

    text = "\n".join(lines)
    if not text.strip():
        raise ParseError("DOCX tidak berisi teks yang bisa dibaca.")

    result = ParsedDocument(pages_total=0, pages_digital=0)
    result.chunks = chunker.chunk_pages([(None, text)])
    return result
