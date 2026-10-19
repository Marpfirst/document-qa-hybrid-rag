"""Generate real PDF/DOCX bytes at test time.

Building the fixtures programmatically instead of committing binaries keeps the
repo clean and makes the *shape* of each case explicit — especially the scanned
and broken-text-layer PDFs, whose whole point is what they lack.
"""

from __future__ import annotations

import io

import fitz
from docx import Document

SOP_SECTIONS = [
    ("SOP CUTI KARYAWAN", 1),
    ("4. Cuti", 2),
    ("4.2 Cuti Tahunan", 3),
    ("Cuti tahunan diberikan sebanyak 12 hari kerja per tahun bagi karyawan "
     "yang telah bekerja sekurang-kurangnya 12 bulan secara terus menerus.", 0),
    ("4.3 Persetujuan", 3),
    ("Pengajuan cuti disetujui oleh atasan langsung paling lambat 7 hari "
     "sebelum tanggal pelaksanaan cuti.", 0),
]


def digital_pdf() -> bytes:
    """A normal typeset PDF with a real text layer and size-based headings."""
    doc = fitz.open()
    page = doc.new_page()
    y = 72.0
    for text, level in SOP_SECTIONS:
        size = {0: 11.0, 1: 20.0, 2: 16.0, 3: 13.0}[level]
        page.insert_textbox(
            fitz.Rect(60, y, 540, y + 90),
            text,
            fontsize=size,
            fontname="helv",
        )
        y += 34 if level else 60
    out = doc.tobytes()
    doc.close()
    return out


def scanned_pdf() -> bytes:
    """A page whose only content is a full-page raster image: no text, no fonts."""
    doc = fitz.open()
    page = doc.new_page()
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 600, 800))
    pix.set_rect(pix.irect, (235, 235, 235))
    page.insert_image(page.rect, pixmap=pix)
    out = doc.tobytes()
    doc.close()
    return out


def mixed_pdf() -> bytes:
    """Scanned cover sheet in front of a digitally typeset body page.

    This is the common real-world shape that per-document routing gets wrong.
    """
    doc = fitz.open()
    cover = doc.new_page()
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 600, 800))
    pix.set_rect(pix.irect, (240, 240, 240))
    cover.insert_image(cover.rect, pixmap=pix)

    body = doc.new_page()
    y = 72.0
    for text, level in SOP_SECTIONS:
        size = {0: 11.0, 1: 20.0, 2: 16.0, 3: 13.0}[level]
        body.insert_textbox(
            fitz.Rect(60, y, 540, y + 90), text, fontsize=size, fontname="helv"
        )
        y += 34 if level else 60
    out = doc.tobytes()
    doc.close()
    return out


def broken_text_layer_pdf() -> bytes:
    """Extracts a plausible amount of text, but with no word breaks —
    the signature of a broken CID map or a bad conversion."""
    doc = fitz.open()
    page = doc.new_page()
    blob = "".join("abcdefghij" for _ in range(40))
    y = 72.0
    for i in range(0, len(blob), 60):
        page.insert_text((60, y), blob[i : i + 60], fontsize=11, fontname="helv")
        y += 14
    out = doc.tobytes()
    doc.close()
    return out


def sop_docx() -> bytes:
    """DOCX using real Word heading styles plus a table."""
    document = Document()
    document.add_heading("SOP CUTI KARYAWAN", level=1)
    document.add_heading("4. Cuti", level=2)
    document.add_heading("4.2 Cuti Tahunan", level=3)
    document.add_paragraph(
        "Cuti tahunan diberikan sebanyak 12 hari kerja per tahun bagi karyawan "
        "yang telah bekerja sekurang-kurangnya 12 bulan secara terus menerus."
    )
    document.add_heading("4.3 Persetujuan", level=3)
    document.add_paragraph(
        "Pengajuan cuti disetujui oleh atasan langsung paling lambat 7 hari "
        "sebelum tanggal pelaksanaan cuti."
    )
    document.add_heading("5. Matriks Kewenangan", level=2)
    table = document.add_table(rows=3, cols=3)
    rows = [
        ["Jenis Cuti", "Penyetuju", "Maksimal Hari"],
        ["Cuti Tahunan", "Atasan Langsung", "12"],
        ["Cuti Besar", "Kepala Divisi", "30"],
    ]
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
