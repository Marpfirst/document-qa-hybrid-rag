"""The evaluation corpus: four Indonesian SOP documents.

The first is the sample SOP that ships with the app (web/public). The other
three are synthetic: written for this evaluation, with made-up rules and made-up
numbers, so the answer to every question can be checked against a known text.
They are built into DOCX bytes at run time, the same way the test fixtures are,
so no binary files are committed.

They are deliberately confusable. "7 hari kerja" appears in the leave SOP and
"5 hari kerja" in the travel SOP; "14 hari" appears for MFA grace, for the
travel advance, and for mandatory training (in working days). A retriever that
only matches on surface words will mix them up.
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document

SAMPLE_SOP = Path(__file__).resolve().parents[2] / "web" / "public" / "contoh-sop-cuti.docx"

# (heading level, text). Level 0 is a body paragraph.
Doc = list[tuple[int, str]]

TRAVEL: Doc = [
    (1, "SOP PERJALANAN DINAS DAN PENGGANTIAN BIAYA"),
    (0, "Dokumen: SOP-FIN-021 | Revisi: 2 | Berlaku: 1 Februari 2026"),
    (2, "1. Tujuan"),
    (0, "Mengatur pengajuan, pelaksanaan, dan penggantian biaya perjalanan dinas agar biaya perusahaan terkendali dan terdokumentasi."),
    (2, "2. Ruang Lingkup"),
    (0, "Berlaku bagi seluruh karyawan tetap yang ditugaskan di luar kota domisili kantor. Perjalanan yang dibiayai pihak ketiga tidak diatur dalam SOP ini."),
    (2, "3. Pengajuan"),
    (3, "3.1 Surat Tugas"),
    (0, "Perjalanan dinas wajib didahului surat tugas yang ditandatangani Kepala Divisi. Pengajuan dilakukan melalui sistem Travel paling lambat 5 hari kerja sebelum keberangkatan."),
    (3, "3.2 Uang Muka"),
    (0, "Uang muka perjalanan dinas dapat diminta sebesar-besarnya 70% dari estimasi biaya dan harus dipertanggungjawabkan paling lambat 14 hari kalender setelah karyawan kembali."),
    (2, "4. Standar Biaya"),
    (3, "4.1 Transportasi"),
    (0, "Tiket pesawat dibeli kelas ekonomi untuk penerbangan di bawah 4 jam. Kelas bisnis hanya untuk golongan IV ke atas pada penerbangan di atas 4 jam."),
    (3, "4.2 Akomodasi"),
    (0, "Batas biaya hotel adalah Rp800.000 per malam untuk golongan III ke bawah dan Rp1.200.000 per malam untuk golongan IV ke atas, sudah termasuk pajak."),
    (3, "4.3 Uang Makan"),
    (0, "Uang makan perjalanan dalam negeri Rp150.000 per hari. Untuk perjalanan luar negeri, uang makan US$50 per hari."),
    (2, "5. Penggantian Biaya"),
    (0, "Klaim penggantian biaya diajukan dengan melampirkan bukti asli melalui sistem Finance paling lambat 10 hari kerja setelah karyawan kembali. Klaim yang melewati batas tersebut ditolak."),
    (2, "6. Pembatalan Perjalanan"),
    (0, "Jika perjalanan dibatalkan, karyawan wajib mengembalikan uang muka dalam 3 hari kerja dan melaporkan pembatalan tiket kepada tim Finance."),
]

SECURITY: Doc = [
    (1, "SOP KEAMANAN INFORMASI DAN PENGGUNAAN PERANGKAT"),
    (0, "Dokumen: SOP-IT-007 | Revisi: 4 | Berlaku: 1 Maret 2026"),
    (2, "1. Tujuan"),
    (0, "Melindungi data dan sistem perusahaan dari akses tidak sah, kebocoran, dan kehilangan."),
    (2, "2. Ruang Lingkup"),
    (0, "Berlaku bagi semua pengguna sistem informasi perusahaan, termasuk karyawan kontrak dan peserta magang."),
    (2, "3. Akun dan Akses"),
    (3, "3.1 Kata Sandi"),
    (0, "Kata sandi minimal 12 karakter dengan kombinasi huruf besar, huruf kecil, angka, dan simbol. Kata sandi wajib diganti setiap 90 hari dan tidak boleh sama dengan 5 kata sandi terakhir."),
    (3, "3.2 Autentikasi Dua Faktor"),
    (0, "Autentikasi dua faktor (MFA) wajib untuk akses email, VPN, dan sistem keuangan. Akun tanpa MFA dinonaktifkan setelah masa tenggang 14 hari."),
    (3, "3.3 Akses Jarak Jauh"),
    (0, "Akses ke jaringan internal dari luar kantor hanya boleh melalui VPN perusahaan. Penggunaan Wi-Fi publik tanpa VPN dilarang."),
    (2, "4. Perangkat"),
    (3, "4.1 Perangkat Milik Perusahaan"),
    (0, "Laptop perusahaan wajib terpasang enkripsi disk dan antivirus yang diperbarui otomatis. Instalasi perangkat lunak di luar daftar resmi memerlukan persetujuan tim IT."),
    (3, "4.2 Media Penyimpanan Eksternal"),
    (0, "Penggunaan flashdisk dan hard disk eksternal dilarang kecuali telah disetujui tertulis oleh Manajer IT."),
    (3, "4.3 Kehilangan Perangkat"),
    (0, "Kehilangan perangkat wajib dilaporkan kepada tim IT dalam 1x24 jam dan, bila perangkat berisi data rahasia, juga kepada kepolisian."),
    (2, "5. Data"),
    (3, "5.1 Klasifikasi Data"),
    (0, "Data diklasifikasikan menjadi empat tingkat: Publik, Internal, Rahasia, dan Sangat Rahasia. Data Rahasia dan Sangat Rahasia tidak boleh dikirim melalui layanan email pribadi."),
    (3, "5.2 Cadangan Data"),
    (0, "Cadangan data sistem dijalankan setiap hari pukul 01.00 dan disimpan selama 30 hari. Uji pemulihan dilakukan setiap 6 bulan."),
    (2, "6. Insiden Keamanan"),
    (0, "Setiap dugaan insiden keamanan, seperti email phishing yang diklik atau akses tidak sah, wajib dilaporkan ke tim SOC melalui hotline keamanan dalam 2 jam sejak diketahui."),
]

ONBOARDING: Doc = [
    (1, "SOP ORIENTASI DAN MASA PERCOBAAN KARYAWAN BARU"),
    (0, "Dokumen: SOP-HR-020 | Revisi: 1 | Berlaku: 1 Januari 2026"),
    (2, "1. Tujuan"),
    (0, "Memastikan karyawan baru memahami perusahaan, memperoleh perangkat kerja, dan dinilai secara adil selama masa percobaan."),
    (2, "2. Ruang Lingkup"),
    (0, "Berlaku bagi karyawan baru berstatus percobaan. Peserta magang tidak termasuk."),
    (2, "3. Hari Pertama"),
    (3, "3.1 Orientasi"),
    (0, "Orientasi dimulai pukul 09.00 di ruang pelatihan lantai 3 dan dipandu staf HR. Karyawan membawa KTP, NPWP, dan buku rekening."),
    (3, "3.2 Perangkat Kerja"),
    (0, "Laptop, akses email, dan kartu akses gedung diserahkan pada hari pertama setelah karyawan menandatangani serah terima aset."),
    (2, "4. Masa Percobaan"),
    (3, "4.1 Durasi"),
    (0, "Masa percobaan berlangsung 3 bulan dengan gaji 100% dari gaji yang disepakati. Status tetap ditetapkan setelah evaluasi akhir dinyatakan lulus."),
    (3, "4.2 Evaluasi"),
    (0, "Evaluasi kinerja dilakukan dua kali oleh atasan langsung, yaitu pada minggu ke-8 dan minggu ke-12. Nilai minimal untuk lulus adalah 70 dari 100."),
    (3, "4.3 Cuti Selama Percobaan"),
    (0, "Hak cuti tahunan belum timbul selama masa percobaan. Cuti hanya diberikan untuk sakit dengan surat keterangan dokter."),
    (2, "5. Pelatihan Wajib"),
    (0, "Pelatihan kepatuhan dan keamanan informasi wajib diselesaikan dalam 14 hari kerja pertama. Sertifikat penyelesaian diunggah ke HRIS."),
    (2, "6. Mentor"),
    (0, "Setiap karyawan baru didampingi seorang mentor yang ditunjuk Kepala Divisi selama 3 bulan pertama."),
]

SYNTHETIC = {
    "sop-perjalanan-dinas.docx": TRAVEL,
    "sop-keamanan-informasi.docx": SECURITY,
    "sop-orientasi-karyawan.docx": ONBOARDING,
}


def build_docx(spec: Doc) -> bytes:
    document = Document()
    for level, text in spec:
        if level == 0:
            document.add_paragraph(text)
        else:
            document.add_heading(text, level=level)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def documents() -> dict[str, bytes]:
    """File name -> DOCX bytes, in a fixed order."""
    files = {"contoh-sop-cuti.docx": SAMPLE_SOP.read_bytes()}
    for name, spec in SYNTHETIC.items():
        files[name] = build_docx(spec)
    return files
